"""QCOS V1.6 收银流水（/cashier_history）视图测试

覆盖（任务九~十六、三十四）：
    1. 页面 / 接口路由可达（admin）
    2. 玩家流水主表（kind=player）+ 桌局汇总（kind=session）+ 柜台商品（kind=product）
    3. 时间筛选：今日 / 昨日 / 近7天 / 本月 / 自定义
    4. 维度筛选：玩家 / 支付方式 / 四八口 / 通宵 / 操作员 / 桌号
    5. 单笔详情（session/players/products/payments/audit）
    6. legacy 历史数据：显示「历史/未知」「历史导入」，不伪造支付方式
    7. Excel 导出：3 个 sheet（玩家收银流水 / 商品明细 / 桌局汇总）
    8. 汇总合计不含 legacy；实收合计与 GMV 穿透一致

安全：临时库 + 真实结账 API 造数，不触碰真实 qcos.db。
"""

import os
import sys
import io
import tempfile
from datetime import datetime, date, timedelta

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

# ===== 测试专用环境变量（必须在 import config 之前设置，非生产凭据）=====
os.environ.setdefault('QCOS_SECRET_KEY', 'test_secret_key_qcos_unit_test')
os.environ.setdefault('QCOS_ADMIN_PASSWORD', 'test_admin_pw_qcos_v2')

import config
_tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.db')
config.DB_PATH = _tmp.name
config.PAYMENT['provider'] = None

import models
models.init_db()
import app as APP
import revenue
from models import get_db

client = APP.app.test_client()

TODAY = date.today().isoformat()
YESTERDAY = (date.today() - timedelta(days=1)).isoformat()

PASS = 0
FAIL = 0


def ok(name):
    global PASS
    PASS += 1
    print(f'  [PASS] {name}')


def bad(name, msg):
    global FAIL
    FAIL += 1
    print(f'  [FAIL] {name}: {msg}')


def check(name, cond, msg=''):
    if cond:
        ok(name)
    else:
        bad(name, msg)


def login():
    r = client.post('/api/auth/login', json={
        'username': 'admin', 'password': os.environ['QCOS_ADMIN_PASSWORD']})
    assert r.status_code == 200, r.data


def now_iso(hours_ago=2):
    return (datetime.now() - timedelta(hours=hours_ago)).isoformat()


def start_session(machine_id, start, players):
    r = client.post('/api/sessions', json={
        'machine_id': machine_id, 'start_time': start, 'players': players})
    assert r.status_code == 201, r.data
    return r.get_json()['id']


def checkout_player(sid, sp_id, payment='cash', **kw):
    base = {
        'payment_method': payment, 'member_id': None, 'product_total': 0,
        'discount_id': None, 'start_time': None, 'is_overnight': False,
        'manual_discount_type': '', 'manual_discount_value': 0,
    }
    base.update(kw)
    r = client.post(f'/api/sessions/{sid}/players/{sp_id}/checkout', json=base)
    assert r.status_code == 200, r.data
    return r.get_json()


def api(url, **params):
    r = client.get(url, query_string=params)
    assert r.status_code == 200, r.data
    return r.get_json()


def seed():
    """造一套可验证流水：
    A) 单人结账（8口机，现金）→ kind=player
    B) 整台结账（4口机，含柜台商品）→ kind=session + kind=product
    C) legacy 历史导入（直接插库，business_date 为空）→ kind=legacy
    """
    db = get_db()

    # A) 单人结账：machine 1 (8port) start now-2h → cash
    start = now_iso(2)
    sid_a = start_session(1, start, [{'name': '流水甲', 'start_time': start}])
    sp_a = client.get(f'/api/sessions/{sid_a}/preview').get_json()['players'][0]['id']
    res_a = checkout_player(sid_a, sp_a)
    fee_a = round(res_a['grand_total'], 2)

    # B) 整台结账 + 柜台商品：machine 3 (4port) start now-3h
    start_b = now_iso(3)
    sid_b = start_session(3, start_b, [{'name': '流水乙', 'start_time': start_b}])
    # 柜台即时售卖（无 session_player_id → 立即 SETTLED + business_date）
    r = client.post('/api/products/sell', json={
        'session_id': sid_b, 'is_custom': True,
        'custom_name': '测试零食', 'custom_category': 'food', 'price': 8.5, 'quantity': 2})
    assert r.status_code == 201, r.data
    prod_total = round(r.get_json()['total'], 2)  # 17.0
    # 整台结账（现金），商品总额 = 柜台商品
    r = client.post(f'/api/sessions/{sid_b}/close', json={
        'payment_method': 'cash', 'product_total': prod_total, 'discount_id': None})
    assert r.status_code == 200, r.data
    fee_b = round(r.get_json()['final_fee'], 2)   # 3h × 10 = 30
    grand_b = round(r.get_json()['grand_total'], 2)  # 30 + 17 = 47

    # C) legacy：已结账但 business_date 为空（历史导入）
    cur = db.execute(
        "INSERT INTO sessions (machine_id, start_time, status, final_fee, source_id) "
        "VALUES (2, ?, 'closed', 25.0, 'legacy_import_test')",
        [TODAY + 'T10:00:00'])
    db.execute('INSERT INTO session_players (session_id, player_name, status) VALUES (?, ?, "playing")',
               [cur.lastrowid, '历史玩家'])
    db.commit()

    return {'sid_a': sid_a, 'fee_a': fee_a,
            'sid_b': sid_b, 'fee_b': fee_b, 'prod_total': prod_total, 'grand_b': grand_b}


# ===================== 1. 页面/接口可达 =====================

def test_routes():
    print('\n[1] 路由可达')
    r = client.get('/cashier_history')
    check('页面 /cashier_history 200', r.status_code == 200, str(r.status_code))
    check('页面含收银流水标识', '收银流水' in r.get_data(as_text=True))
    d = api('/api/cashier_history', preset='today')
    check('接口返回 status=ok', d.get('status') == 'ok', str(d)[:200])
    check('data 含 rows/session_summary/totals/range',
          all(k in d['data'] for k in ('rows', 'session_summary', 'totals', 'range')))


# ===================== 2. 三类流水行 =====================

def test_row_kinds(S):
    print('\n[2] 三类流水行')
    d = api('/api/cashier_history', preset='today')['data']
    rows = d['rows']
    kinds = [r['kind'] for r in rows]
    check('含 player 行', 'player' in kinds, str(kinds))
    check('含 session 行', 'session' in kinds, str(kinds))
    check('含 product 行', 'product' in kinds, str(kinds))
    check('含 legacy 行', 'legacy' in kinds, str(kinds))

    p = next(r for r in rows if r['kind'] == 'player')
    check('player 行字段完整',
          all(k in p for k in ('player_name', 'payment_method', 'business_date', 'paid_at',
                               'grand_total', 'operator', 'machine_name', 'machine_type')),
          str(list(p.keys())))
    check('player.payment_label=现金', p.get('payment_label') == '现金', str(p.get('payment_label')))
    check('player.business_date=今日', p.get('business_date') == TODAY, str(p.get('business_date')))
    check('player.operator=管理员', p.get('operator') == '管理员', str(p.get('operator')))

    s_row = next(r for r in rows if r['kind'] == 'session')
    check('session 行支付方式=现金', s_row.get('payment_label') == '现金', str(s_row.get('payment_label')))
    check('session 行桌型=4port', s_row.get('machine_type') == '4port', str(s_row.get('machine_type')))

    prod = next(r for r in rows if r['kind'] == 'product')
    check('product 行玩家=柜台', prod.get('player_name') == '柜台', str(prod.get('player_name')))
    check('product 行金额=17', abs((prod.get('grand_total') or 0) - 17.0) < 0.01, str(prod.get('grand_total')))
    check('product 行支付方式随桌=现金', prod.get('payment_label') == '现金', str(prod.get('payment_label')))

    leg = next(r for r in rows if r['kind'] == 'legacy')
    check('legacy.payment_label=历史/未知', leg.get('payment_label') == '历史/未知', str(leg.get('payment_label')))
    check('legacy.operator=历史导入', leg.get('operator') == '历史导入', str(leg.get('operator')))
    check('legacy 不伪造支付方式', leg.get('payment_method') is None, str(leg.get('payment_method')))


# ===================== 3. 汇总与穿透 =====================

def test_totals_and_summary(S):
    print('\n[3] 汇总（不含 legacy）与桌局汇总')
    d = api('/api/cashier_history', preset='today')['data']
    tot = d['totals']
    expect_real = round(S['fee_a'] + S['fee_b'] + S['prod_total'], 2)
    check('实收合计 = 单人 + 整台台费 + 柜台商品（不含 legacy）',
          abs(tot['real_paid'] - expect_real) < 0.01, f"{tot['real_paid']} vs {expect_real}")
    check('合计 count = 3（player+session+product）', tot['count'] == 3, str(tot['count']))
    # 桌局汇总按 session 分组（含 legacy 桌）
    sm = d['session_summary']
    check('桌局汇总含 3 桌（2 真实 + 1 legacy）', len(sm) == 3, f'n={len(sm)}')
    s_b = next(s for s in sm if s['session_id'] == S['sid_b'])
    check('整台桌汇总 product_total = 17（柜台商品并入汇总）',
          abs((s_b.get('product_total') or 0) - S['prod_total']) < 0.01, str(s_b))
    check('整台桌汇总 grand_total = 台费+商品',
          abs((s_b.get('grand_total') or 0) - S['grand_b']) < 0.01, str(s_b.get('grand_total')))


# ===================== 4. 时间筛选 =====================

def test_date_filters(S):
    print('\n[4] 时间筛选')
    check('今日 > 0 行', len(api('/api/cashier_history', preset='today')['data']['rows']) >= 3)
    check('昨日 = 0 行', len(api('/api/cashier_history', preset='yesterday')['data']['rows']) == 0)
    check('近7天 >= 今日行数',
          len(api('/api/cashier_history', preset='7d')['data']['rows']) >=
          len(api('/api/cashier_history', preset='today')['data']['rows']))
    check('本月 >= 今日行数',
          len(api('/api/cashier_history', preset='month')['data']['rows']) >=
          len(api('/api/cashier_history', preset='today')['data']['rows']))
    d = api('/api/cashier_history', preset='custom', custom_start=TODAY, custom_end=TODAY)
    check('自定义=今日 行数一致',
          len(d['data']['rows']) == len(api('/api/cashier_history', preset='today')['data']['rows']))
    d2 = api('/api/cashier_history', preset='custom', custom_start=YESTERDAY, custom_end=YESTERDAY)
    check('自定义=昨日 无行', len(d2['data']['rows']) == 0)
    # legacy 行按 start_time 日期归属显示
    d3 = api('/api/cashier_history', preset='custom', custom_start=TODAY, custom_end=TODAY)
    check('legacy 行显示 business_date=今日', any(r['kind'] == 'legacy' and r['business_date'] == TODAY
                                               for r in d3['data']['rows']))


# ===================== 5. 维度筛选 =====================

def test_dimension_filters(S):
    print('\n[5] 维度筛选')
    d = api('/api/cashier_history', preset='today', player='流水甲')['data']['rows']
    check('玩家筛选=流水甲', len(d) == 1 and d[0]['player_name'] == '流水甲', str([r['player_name'] for r in d]))

    d = api('/api/cashier_history', preset='today', payment='cash')['data']['rows']
    check('支付方式筛选=cash 全部命中', len(d) == 3 and all(r['payment_method'] == 'cash' for r in d),
          str([(r['kind'], r['payment_method']) for r in d]))

    d = api('/api/cashier_history', preset='today', port='8')['data']['rows']
    check('四八口筛选=8口（流水甲 + legacy 桌）',
          len(d) == 2 and all(r['machine_type'] == '8port' for r in d),
          str([(r['kind'], r['machine_type']) for r in d]))
    d = api('/api/cashier_history', preset='today', port='4')['data']['rows']
    check('四八口筛选=4口（整台桌 + 柜台商品）',
          len(d) == 2 and all(r['machine_type'] == '4port' for r in d),
          str([(r['kind'], r['machine_type']) for r in d]))

    d = api('/api/cashier_history', preset='today', overnight='0')['data']['rows']
    check('通宵筛选=否（含 legacy）', len(d) == 4, str(len(d)))

    d = api('/api/cashier_history', preset='today', operator='管理员')['data']['rows']
    check('操作员筛选=管理员', len(d) == 3, str([r.get('operator') for r in d]))
    d = api('/api/cashier_history', preset='today', operator='不存在的人')['data']['rows']
    check('操作员筛选=不存在 → 空', len(d) == 0)

    d = api('/api/cashier_history', preset='today', machine=str(S['sid_b'] and 3))['data']['rows']
    check('桌号筛选=machine3', len(d) == 2, str([(r['kind'], r['machine_id']) for r in d]))


# ===================== 6. 单笔详情 =====================

def test_detail(S):
    print('\n[6] 单笔详情')
    r = client.get('/api/cashier_history/detail', query_string={'session_id': S['sid_b']})
    assert r.status_code == 200, r.data
    dd = r.get_json()
    check('详情含 session/players/products/payments/audit',
          all(k in dd for k in ('session', 'players', 'products', 'payments', 'audit')))
    check('详情 session 状态 closed', dd['session']['status'] == 'closed')
    check('详情 players 非空', len(dd['players']) == 1)
    check('详情 products 含柜台商品', len(dd['products']) == 1 and dd['products'][0]['product_name'] == '测试零食',
          str(dd['products']))
    r404 = client.get('/api/cashier_history/detail', query_string={'session_id': 999999})
    check('不存在的桌局 404', r404.status_code == 404)


# ===================== 7. Excel 导出 =====================

def test_export(S):
    print('\n[7] Excel 导出（3 sheet）')
    r = client.get('/api/cashier_history/export', query_string={'preset': 'today'})
    check('导出 200', r.status_code == 200, str(r.status_code))
    check('导出为 xlsx', r.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.data))
        check('sheet 名 = 玩家收银流水/商品明细/桌局汇总',
              wb.sheetnames == ['玩家收银流水', '商品明细', '桌局汇总'], str(wb.sheetnames))
        ws1 = wb['玩家收银流水']
        rows1 = [x for x in ws1.iter_rows(min_row=2, values_only=True) if x and x[0] is not None]
        check('玩家流水 sheet 有 3 行（不含 legacy）', len(rows1) == 3, f'n={len(rows1)}')
        ws2 = wb['商品明细']
        rows2 = [x for x in ws2.iter_rows(min_row=2, values_only=True) if x and x[0] is not None]
        check('商品明细 sheet 有 1 行', len(rows2) == 1, f'n={len(rows2)}')
        ws3 = wb['桌局汇总']
        rows3 = [x for x in ws3.iter_rows(min_row=2, values_only=True) if x and x[0] is not None]
        check('桌局汇总 sheet 有 3 桌（含 legacy）', len(rows3) == 3, f'n={len(rows3)}')
    except Exception as ex:
        bad('Excel 解析', str(ex))


# ===================== 8. 分析中心穿透 =====================

def test_analytics_penetration(S):
    print('\n[8] 分析中心穿透（today_gmv == 流水实收合计）')
    import analytics as A
    dash = A.get_dashboard(get_db(), 'today')
    cashier_total = api('/api/cashier_history', preset='today')['data']['totals']['real_paid']
    check('today_gmv == 流水实收合计（¥994 场景同机制）',
          abs(round(dash['business']['today_gmv'], 2) - cashier_total) < 0.01,
          f"{dash['business']['today_gmv']} vs {cashier_total}")


if __name__ == '__main__':
    print('=' * 60)
    print('QCOS V1.6 收银流水测试')
    print('=' * 60)
    login()
    S = seed()
    test_routes()
    test_row_kinds(S)
    test_totals_and_summary(S)
    test_date_filters(S)
    test_dimension_filters(S)
    test_detail(S)
    test_export(S)
    test_analytics_penetration(S)
    print('\n' + '=' * 60)
    print(f'结果：通过 {PASS} / 失败 {FAIL}')
    print('=' * 60)
    get_db().close()
    try:
        if os.path.exists(_tmp.name):
            os.remove(_tmp.name)
    except OSError:
        pass
    sys.exit(1 if FAIL else 0)
